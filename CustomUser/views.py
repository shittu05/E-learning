from django.shortcuts import render, redirect
from django.contrib.auth.forms import AuthenticationForm
from .forms import RegForm
from django.contrib.auth.models import auth
from django.conf import settings
from django.contrib.auth import get_user_model
User = get_user_model()
from django.contrib.auth import login, authenticate
from django.contrib import messages
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.contrib.admin.views.decorators import staff_member_required
import json
from Myapp.models import Submission


# ──────────────────────────────────────────────────────────────────────────────
# EXCEL HELPER UTILITIES
# ──────────────────────────────────────────────────────────────────────────────

def _apply_header_style(ws, headers, fill_color="4472C4"):
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')


def _auto_column_widths(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _write_user_info(ws, row, submission):
    ws.cell(row=row, column=1, value=submission.user.id)
    ws.cell(row=row, column=2, value=submission.user.name)
    ws.cell(row=row, column=3, value=submission.user.email)


def _excel_response(wb, filename):
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={filename}'
    wb.save(response)
    return response


VIDEO_CORRECT_ANSWERS = {
    'Video One':   {'q1': 'C', 'q2': 'B', 'q3': 'B', 'q4': 'C', 'q5': 'C'},
    'Video Two':   {'q1': 'C', 'q2': 'B', 'q3': 'C', 'q4': 'C', 'q5': 'C'},
    'Video Three': {'q1': 'A', 'q2': 'D', 'q3': 'A', 'q4': 'D', 'q5': 'D'},
    'Video Four':  {'q1': 'A', 'q2': 'B', 'q3': 'C', 'q4': 'D', 'q5': 'A'},
}


# ──────────────────────────────────────────────────────────────────────────────
# AUTH VIEWS
# ──────────────────────────────────────────────────────────────────────────────

def index(request):
    try:
        if request.user.is_authenticated:
            return redirect("/learn/courses/")
    except Exception:
        pass
    return render(request, 'registration.html')


def register(request):
    if request.user.is_authenticated:
        return redirect("/learn/courses/")

    if request.method == 'POST':
        password = request.POST['password']
        email = request.POST['email']
        name = request.POST['name']
        if User.objects.filter(email=email).exists():
            messages.error(request, 'Email already taken')
            return redirect('/register/')
        else:
            user = User.objects.create_user(email=email, password=password, name=name)
            user.save()
            user = authenticate(email=email, password=password)
            if user is not None:
                auth.login(request, user)
                return redirect('/learn/courses/')
    else:
        return render(request, 'registration.html')


def login(request):
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']
        user = authenticate(request, username=username, password=password)
        if user is not None:
            auth.login(request, user)
            return redirect('/learn/courses/')
        else:
            messages.error(request, 'invalid login details')
            form = AuthenticationForm()
            return render(request, 'signin.html', {'form': form})
    else:
        form = AuthenticationForm()
        return render(request, 'signin.html', {'form': form})


def logout(request):
    request.session.clear()
    messages.info(request, "You have been logged out")
    return redirect('login')


# ──────────────────────────────────────────────────────────────────────────────
# EXPORT VIEWS
# ──────────────────────────────────────────────────────────────────────────────

@staff_member_required
def export_page(request):
    return render(request, 'export.html')


@staff_member_required
def export_survey_one(request):
    submissions = Submission.objects.filter(title='Survey One').select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey One"
    _apply_header_style(ws, ['User ID', 'Name', 'Email', 'Answer', 'Retries', 'Submitted At'])
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        ws.cell(row=row, column=4, value=s.answers.get('survey1', 'N/A'))
        ws.cell(row=row, column=5, value=s.retries)                          # ← model field
        ws.cell(row=row, column=6, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)
    return _excel_response(wb, 'survey_one_export.xlsx')


@staff_member_required
def export_survey_two(request):
    submissions = Submission.objects.filter(title='Survey Two').select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey Two"
    _apply_header_style(ws, ['User ID', 'Name', 'Email'] + [f'Q{i}' for i in range(1, 17)] + ['Retries', 'Submitted At'])
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        for i in range(1, 17):
            ws.cell(row=row, column=3 + i, value=s.answers.get(f'question{i}', 'N/A'))
        ws.cell(row=row, column=20, value=s.retries)                         # ← model field
        ws.cell(row=row, column=21, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)
    return _excel_response(wb, 'survey_two_export.xlsx')


def _build_video_quiz_sheet(ws, title, submissions):
    correct_answers = VIDEO_CORRECT_ANSWERS[title]
    _apply_header_style(ws, [
        'User ID', 'Name', 'Email',
        'Q1', 'Q2', 'Q3', 'Q4', 'Q5',
        'Correct Answers', 'Score (%)', 'Retries', 'Submitted At',
    ], fill_color="70AD47")
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        user_answers = {}
        for i in range(1, 6):
            val = s.answers.get(f'question{i}', 'N/A')
            ws.cell(row=row, column=3 + i, value=val)
            user_answers[f'q{i}'] = val
        correct_count = sum(1 for q, ans in user_answers.items() if ans == correct_answers.get(q))
        ws.cell(row=row, column=9,  value=correct_count)
        ws.cell(row=row, column=10, value=f"{(correct_count / 5) * 100:.1f}%")
        ws.cell(row=row, column=11, value=s.retries)                         # ← model field
        ws.cell(row=row, column=12, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)


@staff_member_required
def export_video_one(request):
    title = 'Video One'
    submissions = Submission.objects.filter(title=title).select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _build_video_quiz_sheet(ws, title, submissions)
    return _excel_response(wb, 'video_one_export.xlsx')


@staff_member_required
def export_video_two(request):
    title = 'Video Two'
    submissions = Submission.objects.filter(title=title).select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _build_video_quiz_sheet(ws, title, submissions)
    return _excel_response(wb, 'video_two_export.xlsx')


@staff_member_required
def export_video_three(request):
    title = 'Video Three'
    submissions = Submission.objects.filter(title=title).select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _build_video_quiz_sheet(ws, title, submissions)
    return _excel_response(wb, 'video_three_export.xlsx')


@staff_member_required
def export_video_four(request):
    title = 'Video Four'
    submissions = Submission.objects.filter(title=title).select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = title
    _build_video_quiz_sheet(ws, title, submissions)
    return _excel_response(wb, 'video_four_export.xlsx')


@staff_member_required
def export_post_test(request):
    submissions = Submission.objects.filter(title='Post Test').select_related('user')
    wb = Workbook()
    ws = wb.active
    ws.title = "Post Test"
    _apply_header_style(ws, [
        'User ID', 'Name', 'Email',
        'Q1', 'Q2', 'Q3', 'Q4', 'Q5', 'Q6', 'Q7', 'Q8', 'Q9', 'Q10',
        'Correct Answers', 'Score (%)', 'Retries', 'Submitted At',
    ], fill_color="70AD47")
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        for i in range(1, 11):
            ws.cell(row=row, column=3 + i, value=s.answers.get(f'post{i}', 'N/A'))
        correct = s.answers.get('correct', 0)
        ws.cell(row=row, column=14, value=correct)
        ws.cell(row=row, column=15, value=f"{(correct / 10) * 100:.1f}%")
        ws.cell(row=row, column=16, value=s.retries)                         # ← model field
        ws.cell(row=row, column=17, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)
    return _excel_response(wb, 'post_test_export.xlsx')


@staff_member_required
def export_all_submissions(request):
    wb = Workbook()
    wb.remove(wb.active)

    # Survey One
    submissions = Submission.objects.filter(title='Survey One').select_related('user')
    ws = wb.create_sheet(title="Survey One")
    _apply_header_style(ws, ['User ID', 'Name', 'Email', 'Answer', 'Retries', 'Submitted At'])
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        ws.cell(row=row, column=4, value=s.answers.get('survey1', 'N/A'))
        ws.cell(row=row, column=5, value=s.retries)
        ws.cell(row=row, column=6, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)

    # Survey Two
    submissions = Submission.objects.filter(title='Survey Two').select_related('user')
    ws = wb.create_sheet(title="Survey Two")
    _apply_header_style(ws, ['User ID', 'Name', 'Email'] + [f'Q{i}' for i in range(1, 17)] + ['Retries', 'Submitted At'])
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        for i in range(1, 17):
            ws.cell(row=row, column=3 + i, value=s.answers.get(f'question{i}', 'N/A'))
        ws.cell(row=row, column=20, value=s.retries)
        ws.cell(row=row, column=21, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)

    # Video Quizzes
    for title in ['Video One', 'Video Two', 'Video Three', 'Video Four']:
        submissions = Submission.objects.filter(title=title).select_related('user')
        ws = wb.create_sheet(title=title)
        _build_video_quiz_sheet(ws, title, submissions)

    # Post Test
    submissions = Submission.objects.filter(title='Post Test').select_related('user')
    ws = wb.create_sheet(title="Post Test")
    _apply_header_style(ws, [
        'User ID', 'Name', 'Email',
        'Q1', 'Q2', 'Q3', 'Q4', 'Q5', 'Q6', 'Q7', 'Q8', 'Q9', 'Q10',
        'Correct Answers', 'Score (%)', 'Retries', 'Submitted At',
    ], fill_color="70AD47")
    for row, s in enumerate(submissions, 2):
        _write_user_info(ws, row, s)
        for i in range(1, 11):
            ws.cell(row=row, column=3 + i, value=s.answers.get(f'post{i}', 'N/A'))
        correct = s.answers.get('correct', 0)
        ws.cell(row=row, column=14, value=correct)
        ws.cell(row=row, column=15, value=f"{(correct / 10) * 100:.1f}%")
        ws.cell(row=row, column=16, value=s.retries)
        ws.cell(row=row, column=17, value=s.submitted_at.strftime('%Y-%m-%d %H:%M:%S'))
    _auto_column_widths(ws)

    messages.success(request, 'All submissions exported successfully.')
    return _excel_response(wb, 'all_submissions_export.xlsx')